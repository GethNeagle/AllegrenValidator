import os
import json
import shutil
import re
from collections import defaultdict
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
import ttkbootstrap as tb
from tkinter import scrolledtext, messagebox, filedialog
import threading
import time

APP_VERSION = "1.3"
CONFIG_FILE = "allergen_config.json"

class AllergenValidatorApp:
    def __init__(self, root):
        self.root = root
        self.root.title(f"Allergen Validator v{APP_VERSION}")
        self.root.geometry("1100x700")
        self.root.minsize(1000, 600)

        style = tb.Style()
        style.configure("White.TLabel", foreground="white")

        # ---------------- FOLDERS ----------------
        io_frame = tb.Frame(root, padding=10)
        io_frame.grid(row=0, column=0, columnspan=3, sticky="ew", padx=20, pady=10)
        root.columnconfigure(0, weight=1)
        io_frame.columnconfigure(1, weight=1)

        tb.Label(io_frame, text="Input Folder:", style="White.TLabel").grid(row=0, column=0, sticky="w", pady=5)
        input_container = tb.Frame(io_frame)
        input_container.grid(row=0, column=1, sticky="ew", padx=5)
        self.input_entry = tb.Entry(input_container)
        self.input_entry.pack(side="left", fill="x", expand=True)
        self.input_btn = tb.Button(input_container, text="Browse", bootstyle="primary", command=self.browse_input)
        self.input_btn.pack(side="left", padx=5)

        tb.Label(io_frame, text="Output Folder:", style="White.TLabel").grid(row=1, column=0, sticky="w", pady=5)
        output_container = tb.Frame(io_frame)
        output_container.grid(row=1, column=1, sticky="ew", padx=5, pady=5)
        self.output_entry = tb.Entry(output_container)
        self.output_entry.pack(side="left", fill="x", expand=True)
        self.output_btn = tb.Button(output_container, text="Browse", bootstyle="primary", command=self.browse_output)
        self.output_btn.pack(side="left", padx=5)

        # ---------------- JSON EDITORS ----------------
        json_frame = tb.Frame(root, padding=10)
        json_frame.grid(row=1, column=0, columnspan=3, sticky="nsew", padx=20, pady=10)
        root.rowconfigure(1, weight=1)
        json_frame.columnconfigure(0, weight=1)
        json_frame.columnconfigure(1, weight=1)
        json_frame.rowconfigure(0, weight=1)

        # Mandatory Allergens
        mandatory_frame = tb.Labelframe(json_frame, text="Mandatory Allergens (JSON)", bootstyle="primary")
        mandatory_frame.grid(row=0, column=0, sticky="nsew", padx=5, pady=5)
        mandatory_frame.columnconfigure(0, weight=1)
        mandatory_frame.rowconfigure(0, weight=1)
        self.mandatory_text = scrolledtext.ScrolledText(mandatory_frame, wrap="none", font=("Consolas", 11))
        self.mandatory_text.grid(row=0, column=0, sticky="nsew")

        # Exclusions
        exclusions_frame = tb.Labelframe(json_frame, text="Allergen Exclusions (JSON)", bootstyle="primary")
        exclusions_frame.grid(row=0, column=1, sticky="nsew", padx=5, pady=5)
        exclusions_frame.columnconfigure(0, weight=1)
        exclusions_frame.rowconfigure(0, weight=1)
        self.exclusions_text = scrolledtext.ScrolledText(exclusions_frame, wrap="none", font=("Consolas", 11))
        self.exclusions_text.grid(row=0, column=0, sticky="nsew")

        # ---------------- BUTTONS ----------------
        btn_frame = tb.Frame(root, padding=10)
        btn_frame.grid(row=2, column=0, columnspan=3, pady=5)
        self.run_btn = tb.Button(btn_frame, text="Run Validation", bootstyle="success", command=self.start_validation, width=20)
        self.run_btn.pack(side="left", padx=10)
        self.save_btn = tb.Button(btn_frame, text="Save Allergen Settings", bootstyle="primary", command=self.save_config)
        self.save_btn.pack(side="left", padx=10)
        self.save_email_btn = tb.Button(btn_frame, text="Email Template", bootstyle="primary", command=self.save_email_template)
        self.save_email_btn.pack(side="left", padx=10)

        # ----------------- EMAIL TEMPLATE (hidden) ---------
        self.email_text = scrolledtext.ScrolledText(root, wrap="word", font=("Consolas", 11), height=10)
        self.email_text.pack_forget()
        self.email_text.insert("1.0", self.load_email_template())

        # ---------------- PROGRESS ----------------
        self.progress = tb.Progressbar(root, bootstyle="info", orient="horizontal", length=800, mode="determinate")
        self.progress.grid(row=3, column=0, columnspan=3, pady=10, sticky="ew")
        self.progress_label = tb.Label(root, text="", bootstyle="info")
        self.progress_label.grid(row=4, column=0, columnspan=3)

        # ---------------- VERSION ----------------
        tb.Label(root, text=f"Version {APP_VERSION}", bootstyle="secondary").grid(row=5, column=0, columnspan=3, pady=5)

        # ---------------- LOAD CONFIG ----------------
        config = self.load_config()
        self.mandatory_text.insert("1.0", json.dumps(config["mandatory"], indent=2))
        self.exclusions_text.insert("1.0", json.dumps(config["exclusions"], indent=2))

        # Spinner
        self.spinner_chars = "|/-\\"
        self.spinner_index = 0
        self.current_file_name = ""
        self.spinner_stop_event = None

    # ---------------- CONFIG ----------------
    def load_config(self):
        if os.path.exists(CONFIG_FILE):
            with open(CONFIG_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        return {"mandatory": self.mandatory_allergens(), "exclusions": self.allergen_exclusions()}

    def save_config(self):
        try:
            data = {
                "mandatory": json.loads(self.mandatory_text.get("1.0", "end")),
                "exclusions": json.loads(self.exclusions_text.get("1.0", "end"))
            }
            with open(CONFIG_FILE, "w", encoding="utf-8") as f:
                json.dump(data, f, indent=2)
            messagebox.showinfo("Saved", "Allergen settings saved.")
        except json.JSONDecodeError as e:
            messagebox.showerror("JSON Error", str(e))

    # ----------------- EMAIL TEMPLATE -----------------
    def load_email_template(self):
        template_path = "email_template.txt"
        if os.path.exists(template_path):
            with open(template_path, "r", encoding="utf-8") as f:
                return f.read()
        return """We are currently reviewing the allergen information provided for your products and have identified discrepancies that require clarification.

Our review identified {TOTAL_ISSUES} allergen declaration issue(s).

{MISSING_INGREDIENTS}

Summary of findings:
{SUMMARY}

These issues have been highlighted within the attached spreadsheet for ease of reference.

Kindly review and provide updated ingredient specifications where discrepancies have been identified.

Many thanks,
"""

    def save_email_template(self):
        template_path = "email_template.txt"
        with open(template_path, "w", encoding="utf-8") as f:
            f.write(self.email_text.get("1.0", "end").strip())
        os.system(f'notepad "{template_path}"')

    # ---------------- BROWSERS ----------------
    def browse_input(self):
        folder = filedialog.askdirectory(title="Select Input Folder")
        if folder:
            self.input_entry.delete(0, "end")
            self.input_entry.insert(0, folder)

    def browse_output(self):
        folder = filedialog.askdirectory(title="Select Output Folder")
        if folder:
            self.output_entry.delete(0, "end")
            self.output_entry.insert(0, folder)

                # ---------------- DEFAULT DATA ----------------
    def mandatory_allergens(self):
        return {
            "Contains Sulphur Dioxide/Sulphites": ["sulphur","sulfite","sulphite","sulphites","metabisulphates","metabisulphite","metabisulphites"],
            "Contains Sulphur Dioxide (if greater than 10mg/kg)": ["sulphur","sulfite","sulphite","sulphites","metabisulphite","metabisulphites"],
            "Contains Fish or Fish Products": ["fish"],
            "Contains Crustaceans": ["crustacean","crustaceans"],
            "Contains Molluscs": ["molluscs","mollusc"],
            "Contains Milk or Milk Products": ["milk","milkpowder"],
            "Contains Tree Nuts": ["nuts","almonds","almond","brazil,nut","brazil nuts","cashew","cashew,nuts","cashew,nut","cahsews","hazelnuts","hazelnut","hazlenuts","hazlenut","macadamia","pecan","pecans","walnut","walnuts","pistachio","pistachios"],
            "Almonds": ["almonds","almond"],
            "Brazil Nuts": ["brazil,nut","brazil,nuts"],
            "Cashew Nuts": ["cashew","cashew,nut","cashews"],
            "Hazelnuts": ["hazelnut","hazlenuts","hazlenut","hazelnuts"],
            "Macadamia or Queensland Nuts": ["macadamia","macadamia,nut","macadamia,nuts"],
            "Pecan Nuts": ["pecan","pecans"],
            "Pistachios": ["pistachio","pistachios"],
            "Walnuts": ["walnuts","walnut"],
            "Peanuts": ["peanuts","peanut"],
            "Contains Sesame Seed or Sesame Seed Products": ["sesame"],
            "Contains Celery / Celeriac Products": ["celery","celeriac"],
            "Contains Mustard or Mustard Products": ["mustard"],
            "Contains Eggs / Egg Derivatives": ["egg","eggs"],
            "Contains Lupin Flour / Lupin Products": ["lupin","lupins"],
            "Contains Cereals that Contain Gluten": ["wheat","gluten","barley","rye","oats","spelt","kamut","wheatflour"],
            "Wheat": ["wheat","wheatflour"],
            "Spelt (Wheat)": ["spelt"],
            "Kamut (Wheat)": ["kamut"],
            "Rye": ["rye"],
            "Barley": ["barley"],
            "Oats": ["oats","oat","oatmeal"],
            "Contains Soya": ["soya","soy","soybean","soybeans"]
        }

    def allergen_exclusions(self):
        return {
            "Contains Nuts or Nut Trace": ["nutritional,yeast","nutritional,flakes","nut-free","nut,free","nut,oil,substitute","butternut,squash","coconut","coconut,milk","coconut,cream","coconut,oil","nutmeg","chestnut","chestnuts","pine,nuts"],
            "Contains Tree Nuts": ["pine,nuts","chestnut","chestnuts"],
            "Contains Milk or Milk Products": ["coconut,milk","coconut,cream","almond,milk,alternative","oat,milk","soya,milk","cocoa,butter","butter,beans","butternut,squash","almond,butter","cream,of,tartar","rice,milk"],
            "Contains Fish or Fish Products": ["fish,oil,flavouring,(synthetic)"],
            "Contains Soya": ["fully,refined,soybean,oil","soy,oil","soya,oil","refined,soya,oil","refined,soyabean,oil","refined,soya,bean,oil"],
            "Contains Cereals that Contain Gluten": ["gluten,free","gluten-free"],
            "Wheat": ["buckwheat"],
            "Contains Eggs/Egg Derivatives": ["egg,replacer"]
        }

    # ---------------- CLEAN TEXT ----------------
    @staticmethod
    def clean_text(text):
        text = str(text).lower()
        text = re.sub(r"[^a-z0-9]+", ",", text)
        return re.sub(r",+", ",", text).strip(",")

    # ---------------- THREADED VALIDATION ----------------
    def start_validation(self):
        self.run_btn.config(state="disabled")
        self.save_btn.config(state="disabled")
        self.input_btn.config(state="disabled")
        self.output_btn.config(state="disabled")
        self.input_entry.config(state="disabled")
        self.output_entry.config(state="disabled")
        self.save_email_btn.config(state="disabled")
        threading.Thread(target=self.run_validation, daemon=True).start()

    # ---------------- SPINNER ----------------
    def start_spinner(self):
        self.spinner_stop_event = threading.Event()
        threading.Thread(target=self._spinner_loop, daemon=True).start()

    def _spinner_loop(self):
        while not self.spinner_stop_event.is_set():
            self.spinner_index = (self.spinner_index + 1) % len(self.spinner_chars)
            char = self.spinner_chars[self.spinner_index]
            self.progress_label.config(text=f"{char} Processing: {self.current_file_name}")
            time.sleep(0.15)

    def stop_spinner(self):
        if self.spinner_stop_event:
            self.spinner_stop_event.set()
        self.progress_label.config(text="✅ Processing complete")

    # ---------------- VALIDATION ----------------
    def run_validation(self):
        input_folder = self.input_entry.get()
        output_folder = self.output_entry.get()
        if not input_folder or not output_folder:
            messagebox.showerror("Error", "Select input and output folders.")
            self.run_btn.config(state="normal")
            self.save_btn.config(state="normal")
            return

        try:
            mandatory = json.loads(self.mandatory_text.get("1.0", "end"))
            exclusions = json.loads(self.exclusions_text.get("1.0", "end"))
        except json.JSONDecodeError:
            messagebox.showerror("Error", "Invalid JSON configuration.")
            self.run_btn.config(state="normal")
            self.save_btn.config(state="normal")
            return

        files = [f for f in os.listdir(input_folder) if f.lower().endswith(".xlsx")]
        if not files:
            messagebox.showinfo("Info", "No Excel files found in the input folder.")
            self.run_btn.config(state="normal")
            self.save_btn.config(state="normal")
            return

        self.progress["maximum"] = len(files)
        self.progress["value"] = 0

        # Precompile regex patterns
        compiled_mandatory = {a: [re.compile(rf"\b{re.escape(k)}\b") for k in kw] for a, kw in mandatory.items()}
        compiled_exclusions = {a: [re.compile(re.escape(k)) for k in kw] for a, kw in exclusions.items()}

        # Start spinner
        self.current_file_name = ""
        self.start_spinner()
        

        for idx, file_name in enumerate(files, start=1):
            self.current_file_name = file_name
            self.progress["value"] = idx - 1
            self.root.update_idletasks()

            try:
                input_file = os.path.join(input_folder, file_name)
                output_file = os.path.join(output_folder, f"checked_{file_name}")
                shutil.copy(input_file, output_file)

                df = pd.read_excel(input_file, header=1)
                df["Excel_Row"] = df.index + 3
                df["Validation Notes"] = ""

                allergen_issue_counts = defaultdict(int)
                allergen_issues_rows = defaultdict(list)
                missing_ingredient_count = 0

                for row_idx, row in df.iterrows():
                    ingredients_raw = row.get("Ingredient Declaration", "")
                    if pd.isna(ingredients_raw) or not str(ingredients_raw).strip():
                        df.at[row_idx, "Validation Notes"] = "Ingredient list missing"
                        missing_ingredient_count += 1
                        continue

                    ingredients = self.clean_text(ingredients_raw)
                    issues = []

                    # Sulphur special case
                    sulphur_keywords = compiled_mandatory.get("Contains Sulphur Dioxide/Sulphites", [])
                    sulphur_declared = str(row.get("Contains Sulphur Dioxide/Sulphites", "")).upper().strip()
                    clean_ing = ingredients
                    for pat in compiled_exclusions.get("Contains Sulphur Dioxide/Sulphites", []):
                        clean_ing = pat.sub("", clean_ing)
                    sulphur_found = any(pat.search(clean_ing) for pat in sulphur_keywords)

                    for allergen, patterns in compiled_mandatory.items():
                        if allergen == "Contains Sulphur Dioxide (if greater than 10mg/kg)" and sulphur_declared == "Y" and sulphur_found:
                            continue
                        raw_value = row.get(allergen, "")
                        declared = "" if pd.isna(raw_value) else str(raw_value).upper().strip()

                        clean_ing2 = ingredients
                        for pat in compiled_exclusions.get(allergen, []):
                            clean_ing2 = pat.sub("", clean_ing2)

                        found = any(pat.search(clean_ing2) for pat in patterns)

                        if declared == "Y" and not found:
                            issues.append(f"{allergen}: Y but not found")
                            allergen_issue_counts[allergen] += 1
                            allergen_issues_rows[allergen].append(row["Excel_Row"])
                        elif declared in ["N", "M"] and found:
                            issues.append(f"{allergen}: {declared} but found")
                            allergen_issue_counts[allergen] += 1
                            allergen_issues_rows[allergen].append(row["Excel_Row"])
                        spelt_kamut_allergens = ["Spelt (Wheat)", "Kamut (Wheat)"]

                        if allergen in spelt_kamut_allergens:
                            if declared == "":
                                # Record in counts for email, but don't add to Excel notes
                                allergen_issue_counts[allergen] += 1
                                # Optionally track rows too if you want
                                allergen_issues_rows[allergen].append(row["Excel_Row"])
                        else:
                            if declared == "":
                                issues.append(f"{allergen}: not declared")
                                allergen_issue_counts[allergen] += 1
                                allergen_issues_rows[allergen].append(row["Excel_Row"])

                        if issues:
                            df.at[row_idx, "Validation Notes"] = " | ".join(issues)
                            

                # Write notes to Excel and highlight rows
                wb = load_workbook(output_file)
                ws = wb.active
                ws.insert_cols(1)
                ws.cell(row=2, column=1, value="Validation Notes")
                fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
                for _, r in df.iterrows():
                    row_num = int(r["Excel_Row"])
                    note = r["Validation Notes"]
                    ws.cell(row=row_num, column=1, value=note)
                    if note:
                        for c in range(1, ws.max_column + 1):
                            ws.cell(row=row_num, column=c).fill = fill
                wb.save(output_file)

                # Generate email template
                total_issues = sum(allergen_issue_counts.values())
                summary_lines = [f"• {a}: {c} product(s)" for a, c in allergen_issue_counts.items() if c > 0]
                summary_text = "\n".join(summary_lines)
                missing_text = ""
                if missing_ingredient_count:
                    missing_text = f"{missing_ingredient_count} product(s) had no ingredient data, so allergen checks could not be performed."

                template = self.email_text.get("1.0", "end").strip()
                email_body = template.format(TOTAL_ISSUES=total_issues, MISSING_INGREDIENTS=missing_text, SUMMARY=summary_text)

                email_file_path = os.path.join(output_folder, f"EMAIL_TEMPLATE_{os.path.splitext(file_name)[0]}.txt")
                with open(email_file_path, "w", encoding="utf-8") as f:
                    f.write(email_body)

            except Exception as e:
                messagebox.showerror("Error", f"Error processing {file_name}: {e}")

            self.progress["value"] = idx
            self.root.update_idletasks()



        self.progress_label.config(text="✅ Processing complete")

        # Stop spinner at end
        # Reset buttons
        self.stop_spinner()
        messagebox.showinfo("Complete", "All files processed. Check the output folder.")
        self.run_btn.config(state="normal")
        self.save_btn.config(state="normal")
        self.input_btn.config(state="normal")
        self.output_btn.config(state="normal")
        self.input_entry.config(state="normal")
        self.output_entry.config(state="normal")
        self.save_email_btn.config(state="normal")


# ---------------- RUN ----------------
if __name__ == "__main__":
    root = tb.Window(themename="superhero")
    app = AllergenValidatorApp(root)
    root.mainloop()
